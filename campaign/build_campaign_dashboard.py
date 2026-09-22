#!/usr/bin/env python3
"""Build June–August 2026 analysis. Requires XlsxWriter; openpyxl for validation."""
import csv, sys
from pathlib import Path
from collections import defaultdict
sys.path.insert(0, '/tmp/campaign-analysis-libs')
import xlsxwriter
BASE=Path(__file__).resolve().parent
MONTHS=['June','July','August']
SOURCE=BASE/'GulongPH Web _Google ads_Table - Sheet1.csv'
def number(v):
    v=v.strip().replace('₱','').replace(',','').replace(' ','')
    return None if v.lower() in ('','null','nodata') else float(v)
def div(a,b): return a/b if b else None
def change(a,b): return a/b-1 if b else None
rows=[]; totals_reported={}; month=None
with SOURCE.open(encoding='utf-8-sig',newline='') as f:
    original=list(csv.reader(f))
for line,r in enumerate(original,1):
    label=r[0].strip()
    if label.lower() in ('june 2026','july','august','september'):
        month=label.split()[0].title(); continue
    if month not in MONTHS or not label or label=='Campaign': continue
    if label=='OVERALL TOTAL': totals_reported[month]=number(r[2]); continue
    rows.append(dict(month=month,campaign=label,cost=number(r[2]),sessions=number(r[3]),users=number(r[5]),carts=number(r[6]),checkouts=number(r[7]),purchases=number(r[8]),reported_cpa=number(r[10]),budget=r[11],avg_sales=number(r[12]),source_line=line))
keys=['cost','sessions','users','carts','checkouts','purchases']
tot={m:{k:sum(r[k] or 0 for r in rows if r['month']==m) for k in keys} for m in MONTHS}
checks=[]
for m in MONTHS:
    with (BASE/(m.lower()+'.csv')).open(encoding='utf-8-sig',newline='') as f:
        comparison={r['Campaign']:r for r in csv.DictReader(f)}
    mapping=dict(cost='Cost',sessions='Sessions',users='Total users',carts='Add to carts',checkouts='Checkouts',purchases='Purchase')
    subset=[r for r in rows if r['month']==m]
    assert set(comparison)=={r['campaign'] for r in subset}
    for r in subset:
        for k,col in mapping.items(): assert r[k]==number(comparison[r['campaign']][col]),(m,r['campaign'],k)
    checks.append([m,'Monthly export reconciliation','PASS',f'{len(subset)} campaign rows match all six base metrics'])
    checks.append([m,'Reported total vs campaign spend','PASS' if abs(tot[m]['cost']-totals_reported[m])<.005 else 'DIFFERENCE',f"Row sum ₱{tot[m]['cost']:,.2f}; reported ₱{totals_reported[m]:,.2f}; difference ₱{tot[m]['cost']-totals_reported[m]:.2f}"])
notes=[
'Scope: June, July and August 2026 monthly labels; assumed full months (30/31/31 days). Daily records are unavailable to verify coverage.',
'Primary source: combined Google Ads CSV. Separate june.csv, july.csv and august.csv independently reconcile all six base metrics.',
'July campaign costs total ₱95,066.91; the combined sheet reports ₱95,066.89. Analysis uses the campaign-row sum (₱0.02 difference).',
'All growth is (current − previous) / previous. Source % diff columns are not used. Zero or missing baselines produce N/A.',
'CPA = total cost / total recorded purchases, never the average of campaign CPAs. Purchases are reported events, not verified fulfilled orders or unique buyers.',
'Total users are summed campaign-reported users, not deduplicated monthly or quarterly people. Session totals are sums of reported campaign sessions.',
'July BFG Advantage Touring PMax has zero spend and missing activity. Missing raw values stay blank; monthly sums use available values only.',
'Funnel ratios divide aggregate events; they do not track the same users through sequential stages and may include repeated events.',
'No verified attributed revenue, margins, conversion definitions, attribution windows or tracking-change history are supplied. Actual ROAS and profitability cannot be determined.',
'August Average Net Sales exists for five campaigns only and is not defined. It is retained in Raw Data, but not multiplied by purchases or treated as verified revenue.',
'Campaign mix and June/July month lengths differ. Daily rates aid comparison but do not establish campaign causality.',
'Budget notes are source context; several mid-month changes are listed, so current daily budgets are not substitutes for actual monthly spend.',
]
j,a=tot['July'],tot['August']; q={k:sum(t[k] for t in tot.values()) for k in keys}
insights=[
('Overall result',f"June–August spend was ₱{q['cost']:,.2f}, with {q['purchases']:,.0f} recorded purchases; blended CPA was ₱{q['cost']/q['purchases']:,.2f}."),
('August growth',f"August purchases rose {change(a['purchases'],j['purchases']):.1%} (67 → 188) on {change(a['cost'],j['cost']):.1%} more spend. CPA fell {1-(a['cost']/a['purchases'])/(j['cost']/j['purchases']):.1%} to ₱{a['cost']/a['purchases']:,.2f}."),
('Main purchase driver','Michelin PMax delivered 119 August purchases, 63.3% of total, on 23.9% of spend; CPA was ₱239.50. Its +85 purchases explain 70.2% of the net July–August increase.'),
('Traffic quality signal','August sessions grew 708.1%, but purchases/session fell from 0.348% to 0.121%. Lower cost/session still allowed blended CPA to improve; traffic volume alone overstates purchase progress.'),
('BFG efficiency gap','BFG Advantage Touring PMax used 23.9% of August spend for 7.4% of purchases. Its ₱2,034.65 CPA was 8.5× Michelin PMax CPA; purchase/session was 0.024%.'),
('Search tracking review','SRC Michelin spent ₱20,054.15 in July with 69 sessions and zero purchases. Across the quarter, campaigns with observed zero purchases consumed ₱58,241.83 (21.1% of spend). Check tagging, landing pages and attribution.'),
('Funnel signal','Aggregate checkout/cart ratio improved from 6.54% in July to 8.32% in August, while cart/session dropped from 20.81% to 5.09%. These are event ratios, not cohort conversion rates.'),
('Comparability','July–August continuing campaigns (All Brand PMax, SRC All Brands, Michelin PMax) added 91 purchases. Newly observed/resumed campaigns added 32; absent campaigns removed 2, yielding +121 net.'),
]
actions=[
['1','Validate measurement','Audit session growth, campaign tagging, purchase deduplication and attribution consistency before interpreting traffic changes as demand growth.','Tracking definitions and purchase events reconcile with source systems.'],
['2','Prioritize Michelin PMax review','It has the strongest observed August purchase CPA. Evaluate controlled budget tests after checking fulfilled-order quality and margins.','Monitor incremental purchases, CPA and fulfilled net sales during any test.'],
['3','Review BFG Advantage Touring','High spend share and low purchase share warrant an asset, audience and landing-page review. Compare against BFG Traffic (13 purchases; ₱795.95 CPA).','Evaluate comparable conversion quality and CPA after measurement checks.'],
['4','Investigate SRC Michelin','July spend with zero purchases and only 69 sessions warrants a click-to-session and purchase-tracking audit.','Explain the traffic discrepancy and determine whether purchases were missed.'],
['5','Monitor SRC All Brands','August purchases fell 13 → 10 as spend fell 37.7%; CPA improved 19.0%, so lower volume alone is not evidence of poorer efficiency.','Track both purchase volume and CPA with budget context.'],
]
out=BASE/'GulongPH_Google_Ads_June_September_MTD_2026_Analysis.xlsx'
w=xlsxwriter.Workbook(out)
w.set_properties({'title':'GulongPH Google Ads | June–August 2026','comments':'Built from local CSV exports; metrics recalculated.'})
fmt={}
for k,code in [('text',None),('num','#,##0'),('money','"₱"#,##0.00'),('pct','0.00%'),('decimal','0.00')]:
    fmt[k]=w.add_format({'font_name':'Calibri','font_size':11,'valign':'top',**({'num_format':code} if code else {})})
head=w.add_format({'bold':True,'bg_color':'#16324F','font_color':'white','text_wrap':True,'valign':'vcenter'})
title=w.add_format({'bold':True,'font_size':22,'font_color':'#16324F'})
wrap=w.add_format({'text_wrap':True,'valign':'top','font_size':11})
sub=w.add_format({'bold':True,'font_size':12,'font_color':'#007F86','text_wrap':True,'valign':'top'})
def sheet(name,headers=None,widths=None):
    s=w.add_worksheet(name);s.hide_gridlines(2);s.freeze_panes(1,2);s.set_default_row(21)
    s.set_landscape();s.fit_to_pages(1,0)
    if headers:
        s.write_row(0,0,headers,head);s.set_row(0,42)
    for i,width in enumerate(widths or []):s.set_column(i,i,width)
    return s

def val(s,r,c,v,kind='text'):
    if v is None:s.write(r,c,'N/A',fmt['text'])
    else:s.write(r,c,v,fmt[kind])
d=sheet('Executive Dashboard');d.set_column('A:A',3);d.set_column('B:C',23);d.set_column('D:N',12)
d.merge_range('B2:N3','GulongPH | Google Ads Performance',title)
d.merge_range('B4:N4','June–August 2026 • reported campaign results • PHP',sub)
for i,(label,v) in enumerate([('Total spend',q['cost']),('Recorded purchases',q['purchases']),('Blended CPA',q['cost']/q['purchases'])]):
    c=1+i*4;d.merge_range(5,c,5,c+3,label,head);d.merge_range(6,c,7,c+3,v,w.add_format({'bold':True,'font_size':24,'font_color':'#007F86','num_format':'#,##0' if i==1 else '"₱"#,##0.00'}))
for i,(label,body) in enumerate(insights):
    r=10+i*3;d.merge_range(r,1,r+1,2,label,sub);d.merge_range(r,3,r+1,13,body,wrap);d.set_row(r,23);d.set_row(r+1,23)
d.merge_range('B36:N37','Interpretation: purchase efficiency improved, but the large traffic expansion needs validation. Purchases are reported events; revenue and profitability are not established.',wrap)
metrics=[('Spend','cost','money'),('Sessions','sessions','num'),('Campaign users (sum)','users','num'),('Add to carts','carts','num'),('Checkouts','checkouts','num'),('Recorded purchases','purchases','num'),('CPA','cpa','money'),('Cost / session','cps','money'),('Purchases / session','cvr','pct'),('Carts / session','cart_rate','pct'),('Checkouts / cart','checkout_rate','pct'),('Purchases / checkout','purchase_rate','pct'),('Spend / day','daily_cost','money'),('Purchases / day','daily_purchases','decimal')]
for m,t in tot.items():
    t.update(cpa=div(t['cost'],t['purchases']),cps=div(t['cost'],t['sessions']),cvr=div(t['purchases'],t['sessions']),cart_rate=div(t['carts'],t['sessions']),checkout_rate=div(t['checkouts'],t['carts']),purchase_rate=div(t['purchases'],t['checkouts']),daily_cost=t['cost']/(30 if m=='June' else 31),daily_purchases=t['purchases']/(30 if m=='June' else 31))
s=sheet('Monthly Summary',['Metric',*MONTHS,'July vs June','August vs July','Change basis'],[29,19,19,19,19,19,26])
for i,(label,k,kind) in enumerate(metrics,1):
    s.write(i,0,label)
    for c,m in enumerate(MONTHS,1):val(s,i,c,tot[m][k],kind)
    for c,prev,cur in [(4,'June','July'),(5,'July','August')]:
        v=tot[cur][k]-tot[prev][k] if kind=='pct' else change(tot[cur][k],tot[prev][k]);val(s,i,c,v,'pct')
    s.write(i,6,'Percentage points' if kind=='pct' else 'Relative growth')
s.write(17,0,'Ratios use reported events; campaign user totals are not deduplicated.',wrap);s.merge_range(17,0,18,6,'Ratios use reported events; campaign user totals are not deduplicated. Rate deltas are shown as percentage points (e.g., -0.23% means -0.23 pp).',wrap)
headers=['Month','Campaign','Spend','Sessions','Users (campaign)','Add to carts','Checkouts','Purchases','CPA','Cost/session','Purchase/session','Cart/session','Checkout/cart','Purchase/checkout','Spend share','Purchase share']
s=sheet('Campaign Breakdown',headers,[13,43]+[18]*14)
for i,r in enumerate(rows,1):
    t=tot[r['month']]; values=[r['month'],r['campaign'],*[r[k] for k in keys],div(r['cost'],r['purchases']),div(r['cost'],r['sessions']),div(r['purchases'],r['sessions']) if r['purchases'] is not None else None,div(r['carts'],r['sessions']) if r['carts'] is not None else None,div(r['checkouts'],r['carts']) if r['checkouts'] is not None else None,div(r['purchases'],r['checkouts']) if r['purchases'] is not None else None,r['cost']/t['cost'],div(r['purchases'],t['purchases']) if r['purchases'] is not None else None]
    for c,v in enumerate(values):val(s,i,c,v,'text' if c<2 else 'money' if c in (2,8,9) else 'pct' if c>=10 else 'num')
s.autofilter(0,0,len(rows),15);s.conditional_format(1,8,len(rows),8,{'type':'3_color_scale','min_color':'#63BE7B','max_color':'#F8696B'})
s=sheet('July August Drivers',['Campaign','July spend','August spend','Spend change','July purchases','August purchases','Purchase change','Share of net +121','Comparison status'],[43,18,18,18,18,18,18,21,35])
by={(r['month'],r['campaign']):r for r in rows}
for i,name in enumerate(sorted({r['campaign'] for r in rows if r['month'] in ('July','August')}),1):
    p=by.get(('July',name));c=by.get(('August',name));pc=(p or {}).get('cost') or 0;cc=(c or {}).get('cost') or 0;pp=(p or {}).get('purchases') or 0;cp=(c or {}).get('purchases') or 0
    status='Continuing' if pc and cc else 'Newly observed / resumed' if cc else 'Absent in August'
    for col,v in enumerate([name,pc,cc,cc-pc,pp,cp,cp-pp,(cp-pp)/121,status]):val(s,i,col,v,'text' if col in (0,8) else 'money' if col in (1,2,3) else 'pct' if col==7 else 'num')
s.merge_range(i+3,0,i+4,8,'Bridge only: absent or missing campaign purchases contribute zero to reported-total differences. This does not establish a campaign launch, closure, or a causal effect.',wrap)
s=sheet('Actions',['Priority','Focus','Proposed next step','Success check'],[12,31,90,75])
for i,r in enumerate(actions,1):s.write_row(i,0,r,wrap);s.set_row(i,60)
s=sheet('Data Quality',['Period','Check','Result','Detail'],[16,36,18,110])
for i,r in enumerate(checks,1):s.write_row(i,0,r,wrap);s.set_row(i,32)
s.write_row(9,0,['July','All Brand purchase growth','RECALCULATED','4 → 18 is +350.0%; combined source displays 77.78%.'],wrap)
s.write_row(10,0,['August','SRC All Brands session growth','RECALCULATED','5,544 → 3,140 is −43.36%; combined source displays +183.01%.'],wrap)
s.write_row(11,0,['August','Michelin purchase growth','RECALCULATED','34 → 119 is +250.0%; combined source displays 71.43%.'],wrap)
for r in range(9,12):s.set_row(r,35)
s=sheet('Methodology',['Item','Definition / limitation'],[10,145])
for i,n in enumerate(notes,1):s.write(i,0,i);s.write(i,1,n,wrap);s.set_row(i,42)
s=sheet('Raw Data',['Month','Campaign',*keys,'Reported CPA','Budget note','Average Net Sales (undefined)','Source CSV line'],[13,43]+[19]*7+[52,30,18])
for i,r in enumerate(rows,1):
    values=[r['month'],r['campaign'],*[r[k] for k in keys],r['reported_cpa'],r['budget'],r['avg_sales'],r['source_line']]
    for c,v in enumerate(values):
        if v is not None:val(s,i,c,v,'money' if c in (2,8,10) else 'text' if c in (0,1,9) else 'num')
s.autofilter(0,0,len(rows),11)
s=sheet('Source CSV');s.set_column(0,12,22)
for i,r in enumerate(original):s.write_row(i,0,r)
s=sheet('Charts Data',['Month','Spend','Purchases','CPA','Sessions','Purchase/session'],[16]*6)
for i,m in enumerate(MONTHS,1):
    t=tot[m];s.write_row(i,0,[m,t['cost'],t['purchases'],t['cpa'],t['sessions'],t['cvr']])
for i,(col,name) in enumerate([(1,'Monthly spend (PHP)'),(2,'Recorded purchases'),(3,'Cost per purchase (PHP)'),(5,'Purchases / session')]):
    chart=w.add_chart({'type':'column' if col!=5 else 'line'})
    chart.add_series({'name':name,'categories':['Charts Data',1,0,3,0],'values':['Charts Data',1,col,3,col],'fill':{'color':'#007F86'},'line':{'color':'#007F86'},'data_labels':{'value':True,'num_format':'0.00%' if col==5 else '#,##0'}})
    chart.set_title({'name':name});chart.set_legend({'none':True});chart.set_size({'width':610,'height':300});chart.set_y_axis({'num_format':'0.00%' if col==5 else '#,##0','major_gridlines':{'visible':True}})
    d.insert_chart(39+(i//2)*16,1+(i%2)*7,chart)
d.print_area(0,0,71,14)
s=sheet('Prior Sep Forecast',['Scenario / input','Value / spend','CPA assumption','Projected purchases','Interpretation'],[38,23,23,26,95])
s.write_row(1,0,['Forecast period','September 2026','','','Conditional planning projection, not a statistically calibrated prediction.'],wrap)
s.write(3,0,'Planned September spend');s.write_number(3,1,tot['August']['cost']/31*30,fmt['money'])
s.write(4,0,'Optimistic CPA reduction');s.write_number(4,1,.20,fmt['pct'])
s.write(5,0,'August actual CPA');s.write_number(5,1,tot['August']['cpa'],fmt['money'])
s.write(6,0,'July actual CPA');s.write_number(6,1,tot['July']['cpa'],fmt['money'])
s.write(7,0,'July–August pooled CPA');s.write_number(7,1,(j['cost']+a['cost'])/(j['purchases']+a['purchases']),fmt['money'])
s.merge_range('E4:H5','Editable assumptions: B4 = planned spend; B5 = optimistic CPA reduction. Default spend maintains August daily spend over September’s 30 days.',wrap)
budget=tot['August']['cost']/31*30
scenarios=[('Conservative: July efficiency',tot['July']['cpa'],'=$B$7','Uses July observed CPA; this is a historical stress case, not a lower confidence bound.'),('Pooled recent efficiency',(j['cost']+a['cost'])/(j['purchases']+a['purchases']),'=$B$8','Two-month pooled spend / purchases smooths the most recent improvement.'),('Base: August efficiency',tot['August']['cpa'],'=$B$6','Assumes August campaign mix and measured efficiency continue; default planning baseline.'),('Optimistic: assumed improvement',tot['August']['cpa']*.8,'=$B$6*(1-$B$5)','Assumes 20% lower CPA than August. User-editable scenario, not a fitted trend.')]
for r,(label,cpa,formula,description) in enumerate(scenarios,10):
    s.write(r,0,label);s.write_formula(r,1,'=$B$4',fmt['money'],budget);s.write_formula(r,2,formula,fmt['money'],cpa)
    s.write_formula(r,3,f'=IFERROR(B{r+1}/C{r+1},"N/A")',fmt['decimal'],budget/cpa);s.write(r,4,description,wrap);s.set_row(r,43)
s.merge_range('A17:E18','Statistical limits: n = 3 aggregate months. No daily variation, seasonality history or stable campaign composition is available. Forecasts assume CPA can be held at the scenario level as spend changes; saturation and auction effects are unknown. No confidence interval or validated accuracy claim is presented.',wrap)
s.merge_range('A20:E21','Backtest diagnostic: June CPA applied to July spend predicts 13.89 purchases vs 67 observed (−79.3% error). July CPA applied to August spend predicts 83.95 vs 188 (−55.3%). These two one-step checks show historical CPA carry-forward underpredicted recent growth; they do not validate any forecast model.',wrap)
s.write_row(23,0,['Spend sensitivity','CPA +20%','August CPA','CPA −20%'],head)
for r,factor in enumerate([.8,1,1.2],24):
    s.write_formula(r,0,f'=$B$4*{factor}',fmt['money'],budget*factor)
    for c,cpaf in enumerate([1.2,1,.8],1):s.write_formula(r,c,f'=IFERROR(A{r+1}/($B$6*{cpaf}),"N/A")',fmt['decimal'],budget*factor/(a['cpa']*cpaf))
s.merge_range('A29:E30','Sensitivity cells show expected recorded purchases. They are conditional arithmetic, not probabilities. Compare actual September spend, purchases and fulfilled orders weekly; revise the assumptions when the campaign mix or tracking changes.',wrap)
chart=w.add_chart({'type':'bar'})
chart.add_series({'categories':['Prior Sep Forecast',10,0,13,0],'values':['Prior Sep Forecast',10,3,13,3],'fill':{'color':'#007F86'},'data_labels':{'value':True,'num_format':'0'}})
chart.set_title({'name':'September purchases by scenario'});chart.set_legend({'none':True});chart.set_size({'width':850,'height':330});s.insert_chart('A33',chart)
from september_update import add_september
add_september(globals())
w.close()
report=['# GulongPH Google Ads — June–August 2026','', '| Metric | June | July | August |','|---|---:|---:|---:|']
for label,k,kind in metrics:
    vals=[('₱'+format(tot[m][k],',.2f')) if kind=='money' else format(tot[m][k],'.2%') if kind=='pct' else format(tot[m][k],',.2f' if kind=='decimal' else ',.0f') for m in MONTHS]
    report.append('| '+ ' | '.join([label,*vals])+' |')
report+=['','## Findings','']+[f'**{label}.** {body}\n' for label,body in insights]
report+=['## Next steps','']+[f'{r[0]}. **{r[1]}:** {r[2]}' for r in actions]
report+=['','## September conditional forecast','',f'Default spend: ₱{budget:,.2f}, maintaining August daily spend for 30 days. Conservative July-CPA case: {budget/j["cpa"]:.1f} purchases; pooled July–August case: {budget/((j["cost"]+a["cost"])/(j["purchases"]+a["purchases"])):.1f}; August-efficiency base case: {budget/a["cpa"]:.1f}; optimistic assumed 20% CPA improvement: {budget/(a["cpa"]*.8):.1f}. These are conditional scenarios, not confidence bounds. Three months are insufficient for a validated seasonal forecast. The Excel tab includes editable assumptions, sensitivity analysis and two historical carry-forward checks.']
report+=['','## Source and method','']+['- '+n for n in notes]
(BASE/'GulongPH_Google_Ads_June_September_MTD_2026_Analysis.md').write_text('\n'.join(report)+'\n'+september_text)
from openpyxl import load_workbook
check=load_workbook(out,data_only=True)
assert len(check['Campaign Breakdown']['A'])==17
assert check['Monthly Summary']['D7'].value==188
assert abs(check['Monthly Summary']['C2'].value-95066.91)<.00001
assert len(check['Executive Dashboard']._charts)==4
assert len(rows)==16
print(out)
print('Validated: 16 campaign rows, all monthly CSV metrics reconciled, 4 charts, workbook reopened successfully.')
