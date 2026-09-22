#!/usr/bin/env python3
"""Restyle campaign analysis to match Japan_Promo_SKU_Before_After_Analysis.xlsx."""
import sys,csv
from pathlib import Path
sys.path.insert(0,'/tmp/campaign-analysis-libs')
import xlsxwriter
from openpyxl import load_workbook
B=Path(__file__).resolve().parent
source=B/'GulongPH_Google_Ads_June_September_MTD_2026_Analysis.xlsx'
f=load_workbook(source,data_only=False);v=load_workbook(source,data_only=True)
out=B/'GulongPH_Campaign_Comparison_and_Forecast.xlsx'
w=xlsxwriter.Workbook(out)
title=w.add_format({'bg_color':'#17365D','font_color':'white','font_size':20,'bold':True})
subtitle=w.add_format({'font_size':10,'text_wrap':True,'valign':'vcenter'})
header=w.add_format({'bg_color':'#17365D','font_color':'white','bold':True,'text_wrap':True,'valign':'vcenter'})
section=w.add_format({'bg_color':'#2F75B5','font_color':'white','bold':True})
caution=w.add_format({'bg_color':'#FFF2CC','bold':True,'text_wrap':True,'valign':'vcenter'})
wrap=w.add_format({'text_wrap':True,'valign':'top','font_size':11})
formats={}
def style(code='General',band=False):
    key=(code,band)
    if key not in formats:formats[key]=w.add_format({'num_format':code,'font_size':11,'bg_color':'#F2F4F7' if band else 'white','valign':'top','text_wrap':True})
    return formats[key]
money='"₱"#,##0.00';pct='0.0%';num='#,##0';decimal='0.00'
def page(name,cols=12):
    s=w.add_worksheet(name);s.hide_gridlines(2);s.set_column(0,0,38);s.set_column(1,cols-1,15);s.set_landscape();s.fit_to_pages(1,0);s.freeze_panes(4,1);return s
s=page('Executive Summary')
s.merge_range('A1:L1','GulongPH — Campaign Performance Analysis',title);s.set_row(0,32)
s.merge_range('A2:L2','June–August 2026 completed-month comparison, with September 1–8 actuals and a September forecast.',subtitle)
s.merge_range('A3:L3','PARTIAL MONTH: September includes 8 days only. Compare daily rates and CPA; full-month projections are conditional scenarios, not actual results.',caution);s.set_row(2,34)
for r,c,label,value,code in [(4,0,'August Purchases',188,num),(4,4,'September 1–8 Purchases',50,num),(4,8,'September Forecast • Base',187.5,decimal),(7,0,'August CPA',119119.2/188,money),(7,4,'September 1–8 CPA',38950.37/50,money),(7,8,'CPA Change vs August',(38950.37/50)/(119119.2/188)-1,pct)]:
    s.merge_range(r,c,r,c+2,label,section);s.merge_range(r+1,c,r+1,c+2,value,w.add_format({'bg_color':'#D9EAF7','bold':True,'font_size':16,'num_format':code}));s.set_row(r+1,28)
s.merge_range('A11:L11','Completed-Month Results',section)
s.write_row(11,0,['Metric','June','July','August','July % change','August % change'],header);s.set_row(11,32)
for dest,src in enumerate([2,7,8,3,10,14,15],12):
    row=v['Monthly Summary']; vals=[row.cell(src,c).value for c in range(1,7)]
    for c,value in enumerate(vals):
        code=row.cell(src,c+1).number_format
        s.write(dest,c,value,style(code,dest%2==0))
s.merge_range('A21:L21','Management Answer',section)
findings=[
'August purchases increased 67 → 188 (+180.6%) on 25.3% more spend; blended CPA improved 55.3% to ₱633.61.',
'Michelin PMax drove 119 August purchases (63.3%) and 70.2% of the net July–August purchase increase.',
'September 1–8: ₱38,950.37 spend and 50 purchases; purchases/day increased 3.1% vs August, but CPA worsened 22.9% to ₱779.01.',
'At the observed September pace: approximately 187.5 purchases and ₱146,063.89 spend for the full month. Forecast assumptions are editable.',
'Campaign results are observational. Changes in mix, tracking, weekday composition and conversion lag are uncontrolled; purchases are not verified fulfilled orders.'
]
for r,text in enumerate(findings,21):s.merge_range(r,0,r,11,'• '+text,wrap);s.set_row(r,31)
s.merge_range('A28:F28','Campaign Purchase Changes • July to August',section)
s.write_row(28,0,['Campaign','July purchases','August purchases','Change'],header);s.set_row(28,30)
drivers=[]
for row in v['July August Drivers'].iter_rows(min_row=2,max_row=9,values_only=True):
    if row[0]:drivers.append((row[0],row[4],row[5],row[6]))
for r,row in enumerate(sorted(drivers,key=lambda x:x[3],reverse=True),29):
    for c,value in enumerate(row):s.write(r,c,value,style(num if c else 'General',r%2==0))
    s.set_row(r,30)
chart=w.add_chart({'type':'bar'});chart.add_series({'categories':['Executive Summary',29,0,36,0],'values':['Executive Summary',29,3,36,3],'fill':{'color':'#2F75B5'}});chart.set_title({'name':'Purchase change by campaign'});chart.set_legend({'none':True});chart.set_size({'width':710,'height':350});s.insert_chart('G28',chart)
s.print_area('A1:L46')
# Campaign comparison mirrors the reference workbook's SKU comparison.
data={}
for month,filename in [('June','june.csv'),('July','july.csv'),('August','august.csv'),('September 1–8','sep .csv')]:
    with (B/filename).open(encoding='utf-8-sig',newline='') as h:
        for r in csv.DictReader(h):
            def n(x):
                try:return float(x.replace('₱','').replace(',',''))
                except ValueError:return None
            data[(month,r['Campaign'])]=(n(r['Cost']),n(r['Purchase']))
s=page('Campaign Comparison',13);s.set_column(0,0,46);s.set_column(1,12,19)
s.merge_range('A1:M1','Campaign Comparison',title);s.set_row(0,32)
s.merge_range('A2:M2','Completed monthly purchase counts and CPAs, plus September daily-rate context. N/A means absent or unreported, not a confirmed zero.',subtitle);s.set_row(1,30)
s.write_row(3,0,['Campaign','June purchases','July purchases','August purchases','Jul–Aug change','Jul–Aug % change','Sep 1–8 purchases','Aug purchases/day','Sep purchases/day','Daily % change','August CPA','Sep CPA','CPA % change'],header);s.set_row(3,46)
for i,name in enumerate(sorted({k[1] for k in data}),4):
    counts=[data.get((m,name),(None,None))[1] for m in ['June','July','August','September 1–8']]
    ju,jl,au,se=counts
    ac=data.get(('August',name),(None,None))[0];sc=data.get(('September 1–8',name),(None,None))[0]
    vals=[name,ju,jl,au,au-jl if au is not None and jl is not None else None,au/jl-1 if jl and au is not None else None,se,au/31 if au is not None else None,se/8 if se is not None else None,(se/8)/(au/31)-1 if au and se is not None else None,ac/au if au else None,sc/se if se else None,(sc/se)/(ac/au)-1 if au and se else None]
    for c,value in enumerate(vals):s.write(i,c,'N/A' if value is None else value,style(pct if c in (5,9,12) else money if c in (10,11) else decimal if c in (7,8) else num if c else 'General',i%2==1))
    s.set_row(i,30)
s.autofilter(3,0,i,12)
# Retain all validated analytical tabs, with the reference blue headers and banding.
order=['Monthly Summary','September MTD','September Campaigns','September Forecast','Campaign Breakdown','July August Drivers','Actions','Data Quality','Methodology','Source CSV','Prior Sep Forecast','Charts Data']
for name in order:
    source_sheet=f[name];cached=v[name];s=page(name,source_sheet.max_column)
    for key,dim in source_sheet.column_dimensions.items():
        from openpyxl.utils.cell import column_index_from_string
        c=column_index_from_string(key)-1;s.set_column(c,c,dim.width)
    for r,dim in source_sheet.row_dimensions.items():
        if dim.height:s.set_row(r-1,dim.height)
    merged={str(rng).split(':')[0]:rng for rng in source_sheet.merged_cells.ranges}
    for row in source_sheet:
        for cell in row:
            if cell.value is None:continue
            r,c=cell.row-1,cell.column-1;value=cell.value;form=header if r==0 else style(cell.number_format,r%2==0)
            if cell.coordinate in merged:
                rng=merged[cell.coordinate];s.merge_range(rng.min_row-1,rng.min_col-1,rng.max_row-1,rng.max_col-1,value,caution if 'limit' in str(value).lower() or 'partial' in str(value).lower() else wrap)
            elif cell.data_type=='f':s.write_formula(r,c,value,form,cached.cell(cell.row,cell.column).value)
            else:s.write(r,c,value,form)
    if source_sheet.auto_filter.ref:s.autofilter(source_sheet.auto_filter.ref)
    s.freeze_panes(1,1)
    if name=='September Forecast':
        chart=w.add_chart({'type':'bar'});chart.add_series({'categories':[name,12,0,15,0],'values':[name,12,4,15,4],'fill':{'color':'#2F75B5'},'data_labels':{'value':True,'num_format':'0.0'}});chart.set_title({'name':'September forecast | actuals through September 8'});chart.set_legend({'none':True});chart.set_size({'width':930,'height':330});s.insert_chart('A32',chart)
    if name in ('Source CSV','Charts Data','Prior Sep Forecast'):s.hide()
# All four user-provided monthly exports, preserved verbatim in separate raw tabs.
combined=[]
for month,filename in [('June','june.csv'),('July','july.csv'),('August','august.csv'),('September 1–8','sep .csv')]:
    with (B/filename).open(encoding='utf-8-sig',newline='') as h:
        raw=list(csv.DictReader(h))
    name='Raw September' if month.startswith('September') else 'Raw '+month
    sh=page(name,8);sh.set_column(0,0,46);sh.set_column(1,7,23)
    sh.merge_range('A1:H1',name+' — 2026',title);sh.set_row(0,32)
    sh.merge_range('A2:H2',f'Source: {filename}. '+('September 1–8 only (8 days), per user; incomplete month.' if month.startswith('September') else 'Monthly export; source values retained as supplied.'),caution if month.startswith('September') else subtitle);sh.set_row(1,32)
    sh.write_row(3,0,list(raw[0]),header);sh.set_row(3,32)
    for r,row in enumerate(raw,4):
        for c,value in enumerate(row.values()):sh.write(r,c,value,style('General',r%2==1))
        sh.set_row(r,30)
        combined.append([month,*row.values(),filename,r-2])
    sh.autofilter(3,0,3+len(raw),7)
sh=page('Raw Data',11);sh.set_column(0,0,22);sh.set_column(1,1,46);sh.set_column(2,10,23)
sh.merge_range('A1:K1','Raw Data — June, July, August and September 2026',title);sh.set_row(0,32)
sh.merge_range('A2:K2','22 campaign-month rows from the four monthly CSV files. September is September 1–8 only. Currency strings and No data values are retained exactly as supplied.',caution);sh.set_row(1,34)
sh.write_row(3,0,['Period',*list(raw[0]),'Source file','CSV line'],header);sh.set_row(3,32)
for r,row in enumerate(combined,4):
    for c,value in enumerate(row):sh.write(r,c,value,style('General',r%2==1))
    sh.set_row(r,30)
sh.autofilter(3,0,3+len(combined),10)
assert len(combined)==22
from sep_dec_forecast import add
add(globals())
from updated_summary_forecast import add as add_updated_summary
add_updated_summary(globals())
w.close()
a=load_workbook(out,data_only=True);b=load_workbook(out,data_only=False)
assert a['Executive Summary']['E6'].value==50
assert a['Raw Data'].max_row==26
assert {a['Raw Data'].cell(r,1).value for r in range(5,27)}=={'June','July','August','September 1–8'}
assert a['Raw September'].max_row==10
assert a['September Forecast']['E14'].value==187.5
assert b['September Forecast']['E14'].data_type=='f'
assert not [(sh.title,c.coordinate) for sh in a for row in sh for c in row if c.data_type=='e']
assert len(a['Executive Summary']._charts)==1
print(out)
print('Validated: summary actuals, forecast cached values and editable formulas, charts, no Excel error cells.')
