#!/usr/bin/env python3
"""Update the analytics workbook with blend reconciliation and Michelin gaps."""

import csv
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


BASE = Path(__file__).resolve().parent
WORKBOOK = BASE / "gulong_ad_to_purchase_analytics_ready.xlsx"
BACKUP = BASE / "gulong_ad_to_purchase_analytics_ready_before_reconciliation.xlsx"
RECON = BASE / "gulong_google_ads_blend_reconciliation.csv"

NAVY = "17365D"
BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_RED = "FCE4D6"
WHITE = "FFFFFF"


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def style_data_sheet(ws, table_name: str, widths: dict[str, int]) -> None:
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(str(cell.value), 18)
    table = Table(displayName=table_name, ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def main() -> None:
    rows = read_csv(RECON)
    if not BACKUP.exists():
        shutil.copy2(WORKBOOK, BACKUP)

    wb = load_workbook(WORKBOOK)

    # User-confirmed GA4 metric: Michelin purchases are 69, versus 58 in the
    # decoded request-ID export. Other campaign metrics remain unchanged.
    executive = wb["Executive Summary"]
    executive["A7"] = "GA4 Ecommerce purchases (reported)"
    executive["B7"] = 113
    executive["A32"] = "Reconciliation Update"
    executive["A32"].font = Font(bold=True, color=WHITE)
    executive["A32"].fill = PatternFill("solid", fgColor=NAVY)
    executive.merge_cells("A32:F32")
    executive["A33"] = (
        "Michelin GA4 reports 69 purchase events. The decoded export represents "
        "58 events across 44 unique order IDs; 34 IDs appear in the current blend. "
        "Ten known decoded IDs are missing from the blend, while 11 GA4 events "
        "cannot be identified without a row-level GA4 export."
    )
    executive.merge_cells("A33:F33")
    executive["A33"].alignment = Alignment(wrap_text=True, vertical="top")
    executive.row_dimensions[33].height = 48

    campaign = wb["Campaign Summary"]
    for row in range(2, campaign.max_row + 1):
        if campaign.cell(row, 1).value == "Michelin - PMax PH":
            campaign.cell(row, 10).value = 69
            campaign.cell(row, 11).value = 25
            campaign.cell(row, 10).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            campaign.cell(row, 11).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    how_to = wb["How to Read"]
    existing_terms = {how_to.cell(row, 1).value for row in range(1, how_to.max_row + 1)}
    if "Michelin GA4 Reconciliation" not in existing_terms:
        how_to.append(("Michelin GA4 Reconciliation", "69 reported purchase events = 44 decoded unique IDs plus repeat events and an unresolved 11-event export gap."))
    if "Blend Coverage" not in existing_terms:
        how_to.append(("Blend Coverage", "The current blend contains 40 of 77 decoded IDs, including 34 of 44 Michelin IDs."))

    for sheet_name in ("Order Reconciliation", "Michelin Order IDs", "Michelin Gaps"):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    detail = wb.create_sheet("Order Reconciliation")
    headers = list(rows[0])
    detail.append(headers)
    for item in rows:
        detail.append([item[h] for h in headers])
        if item["blend_match_status"] == "Missing from blend":
            for cell in detail[detail.max_row]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
    style_data_sheet(
        detail,
        "OrderReconciliationTable",
        {"order_id": 12, "decoded_campaign": 34, "blend_match_status": 20,
         "field_match_status": 18, "sku": 58, "purchased_brand": 20,
         "backend_customer_type": 22, "backend_customer_source": 24,
         "net_sales_amount": 18},
    )

    michelin = wb.create_sheet("Michelin Order IDs")
    michelin_rows = [item for item in rows if item["decoded_campaign"] == "Michelin - PMax PH"]
    michelin_headers = [
        "order_id", "ecommerce_purchases", "blend_match_status", "field_match_status",
        "order_status", "payment_status", "net_sales_amount", "quantity", "sku",
        "purchased_brand", "backend_customer_type", "backend_customer_source",
        "customer_source_in_blend",
    ]
    michelin.append(michelin_headers)
    for item in michelin_rows:
        michelin.append([
            item["order_id"], item["ecommerce_purchases"], item["blend_match_status"],
            item["field_match_status"], item["decoded_order_status"],
            item["decoded_payment_status"], item["net_sales_amount"], item["quantity"],
            item["sku"], item["purchased_brand"], item["backend_customer_type"],
            item["backend_customer_source"], item["customer_source_in_blend"],
        ])
        if item["blend_match_status"] == "Missing from blend":
            for cell in michelin[michelin.max_row]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_RED)
    style_data_sheet(
        michelin,
        "MichelinOrderIDsTable",
        {"order_id": 12, "ecommerce_purchases": 22, "blend_match_status": 20,
         "field_match_status": 18, "order_status": 18, "payment_status": 20,
         "net_sales_amount": 18, "quantity": 12, "sku": 58,
         "purchased_brand": 20, "backend_customer_type": 22,
         "backend_customer_source": 24, "customer_source_in_blend": 24},
    )

    gaps = wb.create_sheet("Michelin Gaps")
    gap_headers = [
        "order_id", "gap_type", "ecommerce_purchases", "order_status",
        "payment_status", "net_sales_amount", "purchased_brand", "sku",
        "backend_customer_type", "backend_customer_source"
    ]
    gaps.append(gap_headers)
    missing = [
        item for item in rows
        if item["decoded_campaign"] == "Michelin - PMax PH"
        and item["blend_match_status"] == "Missing from blend"
    ]
    for item in missing:
        gaps.append([
            item["order_id"], "Decoded ID missing from blend", item["ecommerce_purchases"],
            item["decoded_order_status"], item["decoded_payment_status"],
            item["net_sales_amount"], item["purchased_brand"], item["sku"],
            item["backend_customer_type"], item["backend_customer_source"],
        ])
    # The unresolved difference is event-level and has no known order ID yet.
    gaps.append(["", "11 GA4 purchase events not represented in decoded export", 11, "", "", "", "", "", "", ""])
    style_data_sheet(
        gaps,
        "MichelinGapsTable",
        {"order_id": 12, "gap_type": 52, "ecommerce_purchases": 22,
         "order_status": 18, "payment_status": 20, "net_sales_amount": 18,
         "purchased_brand": 20, "sku": 58, "backend_customer_type": 22,
         "backend_customer_source": 24},
    )

    wb.save(WORKBOOK)
    print(f"Updated: {WORKBOOK}")
    print(f"Backup: {BACKUP}")
    print(f"Reconciliation rows: {len(rows)}")
    print(f"Known Michelin blend gaps: {len(missing)}")
    print(f"Updated at: {datetime.now().isoformat(timespec='seconds')}")


if __name__ == "__main__":
    main()
