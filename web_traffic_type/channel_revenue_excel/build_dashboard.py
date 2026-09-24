import sys,csv,json,base64
sys.path.insert(0,'/tmp/channel_excel_libs')
from pathlib import Path
from urllib.parse import urlsplit,parse_qs
from decimal import Decimal
from datetime import datetime
from collections import Counter
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name
B=Path(__file__).resolve().parent
sources={'Google Ads':'revenue_match/google_cpc.csv','Direct & Organic':'direct_and_organic_revenue/direct_and_organic.csv','Referral':'referral_revenue/referral.csv'}
backend=json.loads((B/'backend_orders.json').read_text()); lookup={r['order_id']:r for r in backend};assert len(lookup)==len(backend)
rows=[];originals={}
for channel,file in sources.items():
 originals[channel]=list(csv.DictReader((B.parent/file).open(encoding='utf-8-sig')))
 for n,r in enumerate(originals[channel],2):
  oid=base64.b64decode(parse_qs(urlsplit(r['Page path + query string']).query)['requestid'][0],validate=True).decode();o=lookup.get(oid,{})
  date=o.get('order_date','')[:10];scope='Included' if '2026-09-01'<=date<='2026-09-24' else ('Outside period' if date else 'Unmatched')
  good=o.get('is_fulfillment_success_order')=='true' and o.get('is_excluded_test_order')=='false'
  net=float(o['net_sales_amount']) if o.get('net_sales_amount') is not None else None
  status=o.get('order_status_raw','Unknown')
  rows.append({'Channel':channel,'Session campaign':r['Session campaign'],'Order ID':oid,'Order date':datetime.fromisoformat(o['order_date']) if date else None,'Week':f'Week {(int(date[-2:])-1)//7+1}' if scope=='Included' else 'Outside / unknown','Scope':scope,'Customer type':o.get('customer_type_raw','Unknown'),'Customer source':o.get('customer_source_raw','Unknown'),'Order status':status,'Payment status':o.get('payment_status_raw','Unknown'),'Net sales':net,'Fulfilled net revenue':net if good else (0 if o else None),'Gross order value':float(o['gross_sales_amount']) if o else None,'Fulfilled flag':int(good),'Cancelled flag':int(status=='Cancelled'),'Open net value':net if status in ['New','Pending','Processing'] else (0 if o else None),'Cancelled net value':net if status=='Cancelled' else (0 if o else None),'Return refund net value':net if status in ['For Return','For Refund'] else (0 if o else None),'Purchase events':int(r['Ecommerce purchases']),'Quantity':int(o['quantity_int']) if o.get('quantity_int') else None,'Brand':o.get('brand',''),'SKU':o.get('sku',''),'Page path + query string':r['Page path + query string'],'Source file':file,'Source CSV row':n,'Backend loaded at':o.get('_ingest_loaded_at',''),'Backend updated at':o.get('date_updated',''),'Test order':o.get('is_excluded_test_order','Unknown')})
assert len({r['Order ID'] for r in rows})==len(rows)
included=[r for r in rows if r['Scope']=='Included'];excluded=[r for r in rows if r['Scope']!='Included']
headers=list(rows[0]); N=len(rows)+1
path=B/'Gulong_Channel_Revenue_September_2026.xlsx'
w=xlsxwriter.Workbook(path,{'strings_to_urls':False,'strings_to_formulas':False})
w.set_properties({'title':'Gulong | September Channel Revenue','subject':'September 1–24, 2026 order cohorts','author':'Gulong Analytics'})
navy='#17324D';teal='#008C95';blue='#3975B8';gold='#E9AE43';red='#C65757'
f={
 'title':w.add_format({'bold':True,'font_size':22,'font_color':'white','bg_color':navy}),
 'sub':w.add_format({'font_color':'#596B7B','font_size':11,'text_wrap':True,'valign':'vcenter'}),
 'head':w.add_format({'bold':True,'font_color':'white','bg_color':navy,'text_wrap':True}),
 'money':w.add_format({'num_format':'"₱"#,##0.00;[Red]("₱"#,##0.00)'}),
 'pct':w.add_format({'num_format':'0.0%'}),'date':w.add_format({'num_format':'mmm d, yyyy hh:mm'}),
 'integer':w.add_format({'num_format':'#,##0'}),'text':w.add_format({'text_wrap':True,'valign':'top'}),
 'card':w.add_format({'bold':True,'font_size':23,'font_color':navy,'bg_color':'#EAF3F7','num_format':'"₱"#,##0'}),
 'label':w.add_format({'bold':True,'font_color':teal,'bg_color':'#EAF3F7'}),
}
def sheet(name):
 s=w.add_worksheet(name);s.hide_gridlines(2);s.set_tab_color(teal);s.set_default_row(20);s.set_landscape();s.fit_to_pages(1,0);return s
def banner(s,title,subtitle,end=10):
 s.merge_range(0,0,1,end,title,f['title']);s.merge_range(2,0,3,end,subtitle,f['sub']);s.set_row(2,23)
def data_sheet(name,data):
 s=sheet(name);hs=list(data[0]) if data else headers
 s.freeze_panes(1,3);s.set_column(0,len(hs)-1,19)
 vals=[]
 for r in data:vals.append([r.get(h) for h in hs])
 cols=[]
 for h in hs:
  fmt=f['money'] if h in ['Net sales','Fulfilled net revenue','Gross order value','Open net value','Cancelled net value','Return refund net value'] else f['date'] if h=='Order date' else None
  cols.append({'header':h,**({'format':fmt} if fmt else {})})
 s.add_table(0,0,len(vals),len(hs)-1,{'data':vals,'columns':cols,'style':'Table Style Medium 2'})
 for i,h in enumerate(hs):
  if h in ['SKU','Page path + query string','Source file']:s.set_column(i,i,45)
 return s
# Dashboard first, then analytic tabs.
dash=sheet('Executive Dashboard');dash.set_column('A:L',13);banner(dash,'GULONG | CHANNEL REVENUE','September 1–24, 2026 • Backend order-date cohorts • Week 4 covers 3 days • Values exclude VAT',11)
raw=data_sheet('Source Data',rows)
for ch,name in [('Google Ads','Original Google Ads'),('Direct & Organic','Original Direct Organic'),('Referral','Original Referral')]:data_sheet(name,originals[ch])
data_sheet('Backend Extract',backend);data_sheet('Exceptions',excluded)
def rng(h):
 c=xl_col_to_name(headers.index(h));return f"'Source Data'!${c}$2:${c}${N}"
def summetric(rs,h):return sum(Decimal(str(r[h] or 0)) for r in rs)
metrics=['Orders','Net sales','Fulfilled orders','Fulfilled net revenue','Fulfillment rate','Cancelled orders','Cancelled net value','Open net value','Return refund net value','Net value per order','Fulfilled value share']
summary_refs={}
def summary(name,keys,groups):
 s=sheet(name);s.set_column(0,len(keys)-1,23);s.set_column(len(keys),len(keys)+len(metrics)-1,21)
 banner(s,name,'September 1–24 only. Order values are counted once; status reflects the backend snapshot.',len(keys)+len(metrics)-1)
 for c,h in enumerate(keys+metrics):s.write(5,c,h,f['head'])
 s.set_row(5,34);s.freeze_panes(6,len(keys));out=[]
 for rowno,g in enumerate(groups,6):
  rs=[r for r in included if all(r[k]==v for k,v in zip(keys,g))]
  criteria=[f'{rng("Scope")},"Included"']+[f'{rng(k)},"{v}"' for k,v in zip(keys,g)]
  criteria=','.join(criteria)
  for c,val in enumerate(g):s.write(rowno,c,val)
  count=len(rs);net=summetric(rs,'Net sales');ful=sum(r['Fulfilled flag'] for r in rs);rev=summetric(rs,'Fulfilled net revenue');can=sum(r['Cancelled flag'] for r in rs)
  vals=[count,net,ful,rev,ful/count if count else 0,can,summetric(rs,'Cancelled net value'),summetric(rs,'Open net value'),summetric(rs,'Return refund net value'),net/count if count else 0,rev/net if net else 0]
  forms=[f'COUNTIFS({criteria})',f'SUMIFS({rng("Net sales")},{criteria})',f'SUMIFS({rng("Fulfilled flag")},{criteria})',f'SUMIFS({rng("Fulfilled net revenue")},{criteria})',None,f'SUMIFS({rng("Cancelled flag")},{criteria})',f'SUMIFS({rng("Cancelled net value")},{criteria})',f'SUMIFS({rng("Open net value")},{criteria})',f'SUMIFS({rng("Return refund net value")},{criteria})',None,None]
  start=len(keys);rr=rowno+1
  for i in [4,9,10]:
   a,b={4:(2,0),9:(1,0),10:(3,1)}[i];forms[i]=f'IFERROR({xl_col_to_name(start+a)}{rr}/{xl_col_to_name(start+b)}{rr},0)'
  for i,(val,formula) in enumerate(zip(vals,forms)):
   fmt=f['pct'] if i in [4,10] else f['integer'] if i in [0,2,5] else f['money']
   s.write_formula(rowno,start+i,'='+formula,fmt,float(val))
  out.append(dict(zip(keys+metrics,list(g)+[float(v) for v in vals])))
 s.autofilter(5,0,5+len(groups),len(keys)+len(metrics)-1)
 s.conditional_format(6,len(keys)+4,5+len(groups),len(keys)+4,{'type':'3_color_scale'})
 s.conditional_format(6,len(keys)+1,5+len(groups),len(keys)+1,{'type':'data_bar','bar_color':teal})
 summary_refs[name]=(s,keys,groups);return out
channels=list(sources);weeks=[f'Week {i}' for i in range(1,5)]
channelstats=summary('Channel Summary',['Channel'],[(c,) for c in channels])
weekly=summary('Weekly Summary',['Week','Channel'],[(wk,c) for wk in weeks for c in channels])
types=sorted({r['Customer type'] for r in included})
typestats=summary('Customer Type',['Customer type'],[(t,) for t in types])
summary('Channel Customer Type',['Channel','Customer type'],[(c,t) for c in channels for t in types])
summary('Status Breakdown',['Order status','Channel'],[(st,c) for st in sorted({r['Order status'] for r in included}) for c in channels])
summary('Campaign Breakdown',['Session campaign'],[(c,) for c in sorted({r['Session campaign'] for r in included})])
# Compact chart source uses formulas linked to weekly summary with cached results.
chdata=sheet('Chart Data');chdata.write_row(0,0,['Week']+channels+['Total net','Total fulfilled','Days','Net per day'])
for i,wk in enumerate(weeks,1):
 chdata.write(i,0,wk+(' (partial)' if i==4 else ''))
 vals=[];revs=[]
 for j,c in enumerate(channels,1):
  record=next(r for r in weekly if r['Week']==wk and r['Channel']==c);v=record['Net sales'];vals.append(v);revs.append(record['Fulfilled net revenue'])
  sr=7+(i-1)*3+j-1;chdata.write_formula(i,j,f"='Weekly Summary'!D{sr}",f['money'],v)
 chdata.write_formula(i,4,f'=SUM(B{i+1}:D{i+1})',f['money'],sum(vals));chdata.write(i,5,sum(revs),f['money']);chdata.write(i,6,3 if i==4 else 7);chdata.write_formula(i,7,f'=E{i+1}/G{i+1}',f['money'],sum(vals)/(3 if i==4 else 7))
chdata.set_column('A:H',22)
def chart(title,kind='column',subtype=None):
 c=w.add_chart({'type':kind,**({'subtype':subtype} if subtype else {})});c.set_title({'name':title,'name_font':{'size':12,'color':navy}});c.set_style(10);c.set_size({'width':550,'height':300});c.set_legend({'position':'bottom'});c.set_y_axis({'num_format':'₱#,##0','major_gridlines':{'visible':True,'line':{'color':'#E6EDF2'}}});return c
c=chart('Weekly net sales by source channel','column','stacked')
for j,(name,color) in enumerate(zip(channels,[blue,teal,gold]),1):c.add_series({'name':name,'categories':['Chart Data',1,0,4,0],'values':['Chart Data',1,j,4,j],'fill':{'color':color}})
dash.insert_chart('A10',c)
c=chart('Order value versus fulfilled revenue')
for j,(name,color) in enumerate([('Net sales',blue),('Fulfilled revenue',teal)]):c.add_series({'name':name,'categories':['Channel Summary',6,0,8,0],'values':['Channel Summary',6,2 if j==0 else 4,8,2 if j==0 else 4],'fill':{'color':color}})
dash.insert_chart('G10',c)
c=chart('Daily net order value | unequal week lengths','line');c.add_series({'name':'Net sales / days','categories':['Chart Data',1,0,4,0],'values':['Chart Data',1,7,4,7],'line':{'color':teal,'width':3},'marker':{'type':'circle'}});dash.insert_chart('A26',c)
c=chart('Net sales by backend customer type','bar');c.add_series({'name':'Net sales','categories':['Customer Type',6,0,5+len(types),0],'values':['Customer Type',6,2,5+len(types),2],'fill':{'color':blue}});dash.insert_chart('G26',c)
# Headline cards.
net=summetric(included,'Net sales');rev=summetric(included,'Fulfilled net revenue');cancel=summetric(included,'Cancelled net value');openval=summetric(included,'Open net value')
for i,(label,value) in enumerate([('ALL-ORDER NET SALES',net),('FULFILLED NET REVENUE',rev),('OPEN ORDER NET VALUE',openval),('CANCELLED NET VALUE',cancel)]):
 col=i*3;dash.merge_range(5,col,5,col+2,label,f['label']);dash.merge_range(6,col,7,col+2,float(value),f['card'])
insights=[]
best=max(channelstats,key=lambda r:r['Fulfilled net revenue'])
insights.append(f"{best['Channel']} contributes {best['Fulfilled net revenue']/float(rev):.1%} of fulfilled net revenue (₱{best['Fulfilled net revenue']:,.0f}).")
insights.append(f"Cancelled orders represent ₱{cancel:,.0f}, or {cancel/net:.1%} of September order value. Review cancellation reasons before increasing acquisition spend.")
insights.append(f"Open orders hold ₱{openval:,.0f}. Prioritize payment and fulfillment follow-up; this value is not a revenue forecast.")
v=[sum(r['Net sales'] for r in weekly if r['Week']==wk) for wk in weeks]
insights.append(f"Week 4 averages ₱{v[3]/3:,.0f} per day versus ₱{v[2]/7:,.0f} in Week 3 ({v[3]/3/(v[2]/7)-1:+.1%}). Three-day coverage and weekday mix limit comparison.")
insights.append('Week 4 has no fulfilled revenue in this snapshot. These recent order cohorts have had less time to fulfill; this is not evidence of zero eventual sales.')
insights.append(f"{sum(r['Payment status']=='Partial Payment' and r['Fulfilled flag']==1 for r in included)} fulfilled orders are partially paid. Fulfilled revenue uses full order value and is not cash collected.")
insights.append('Backend customer type is an operational classification (Website/Fb/Chatbot), not new versus returning customer and not the analytics acquisition channel.')
insights.append('No advertising spend, clicks or session denominators were supplied. ROAS, CAC and visitor-to-order conversion cannot be calculated from these exports.')
dash.merge_range('A43:L43','ANALYST FINDINGS & NEXT ACTIONS',f['head'])
for i,t in enumerate(insights):dash.merge_range(44+i*2,0,45+i*2,11,t,f['text'])
dash.set_print_scale(65);dash.print_area('A1:L61')
notes=sheet('Read Me');notes.set_column('A:A',29);notes.set_column('B:B',115);banner(notes,'READ ME | Scope & methodology','Reproducible analysis with source records and formula-based summaries.',5)
notesdata=[('Period','September 1–24, 2026 using backend order_date, not analytics event date. Dates are backend DATETIME values without timezone conversion.'),('Weeks','Week 1 Sep 1–7; Week 2 Sep 8–14; Week 3 Sep 15–21; Week 4 Sep 22–24 (3 days).'),('Snapshot',str(sorted({r['_ingest_loaded_at'] for r in backend}))),('Source mapping','Google Ads = first supplied export, including its cross-network row. Direct & Organic combines two campaign labels; Campaign Breakdown keeps them separate. Referral = referral.csv.'),('Net sales','Existing gulong_core.orders_raw_deduped calculation: ROUND(total_amount_due / 1.12, 2). Source table is gulong_backend.orders_raw. Cancelled, unfinished, returns and refunds remain in all-order net sales.'),('Fulfilled revenue','is_fulfillment_success_order = true and is_excluded_test_order = false. Full order value is used even for partial payment. This is not cash collected or a formal accounting revenue-recognition schedule.'),('Customer type','customer_type_raw from backend, preserving the backend label. Unknown means no backend match. It does not describe new/returning status.'),('Coverage',f'{len(rows)} CSV records / distinct order IDs; {len(backend)} backend matches; {len(included)} September orders; {len(excluded)} exceptions. No duplicate or overlapping order IDs.'),('Excluded orders','August orders 38959 and 38805, plus unmatched IDs 39900, 39109, 39052. Source Data retains all records; dashboard includes only September 1–24.'),('Event counts','437 purchase events across 267 unique IDs. Purchase events are not revenue multipliers.'),('Summary formulas','COUNTIFS and SUMIFS reference Source Data, with cached results for preview. Changes within existing source rows recalculate in Excel. New rows require rebuilding to extend formula ranges. Analyst findings are snapshot commentary.'),('Source files','Three Original tabs retain the input columns and values. Backend Extract retains the retrieved query result. Source Data contains the joined analysis fields. Exceptions lists excluded/unmatched records.'),('Comparison caution','Weekly figures are order-date cohorts using present order status, not sales fulfilled during each calendar week. Recent cohorts have less time to fulfill.'),('Design reference','japan_sku: executive dashboard, summary and detailed breakdown tabs, source data, peso formatting and performance charts.'),('Refresh','Run match_orders.sql via bq to replace backend_orders.json, then run build_dashboard.py with XlsxWriter available. Files saved in channel_revenue_excel.')]
for i,(a,b) in enumerate(notesdata,5):notes.write(i,0,a,f['head']);notes.write(i,1,b,f['text']);notes.set_row(i,48)
# CSV companion with customer type.
with (B/'source_data_with_customer_type.csv').open('w',encoding='utf-8-sig',newline='') as fp:
 dw=csv.DictWriter(fp,fieldnames=headers);dw.writeheader();dw.writerows(rows)
w.close()
(B/'analysis_summary.json').write_text(json.dumps({'rows':len(rows),'matched':len(backend),'included':len(included),'net':str(net),'fulfilled':str(rev),'channels':channelstats,'customer_types':typestats,'insights':insights},indent=2))
# Independent workbook and reconciliation checks.
import openpyxl
book=openpyxl.load_workbook(path,data_only=True)
assert book['Source Data'].max_row==len(rows)+1
assert abs(book['Channel Summary']['C7'].value-channelstats[0]['Net sales'])<0.001
assert abs(sum(book['Channel Summary'].cell(r,3).value for r in range(7,10))-float(net))<0.001
assert abs(sum(book['Channel Summary'].cell(r,5).value for r in range(7,10))-float(rev))<0.001
assert sum(r['Purchase events'] for r in rows)==437
assert len(book['Executive Dashboard']._charts)==4
assert all(cell.data_type!='e' for sh in book for row in sh for cell in row)
assert len(book['Source Data'].tables)==1
for r in included:
 assert abs(r['Net sales']-round(r['Gross order value']/1.12,2))<0.011
print(json.dumps({'file':str(path),'sheets':book.sheetnames,'rows':len(rows),'included':len(included),'net':str(net),'fulfilled':str(rev),'insights':insights},indent=2))
